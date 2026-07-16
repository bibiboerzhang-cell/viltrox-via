"""V-KPI on-demand CSV/XLSX/PDF exports."""
from __future__ import annotations

import csv
import io
import json
import secrets
from datetime import datetime, timedelta, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.core.logging import get_logger
from app.core.permissions import check_tab_permission
from app.db.connection import get_conn
from app.domains import reports
from app.domains import audit, business_truth
from app.domains.access import scope
from app.domains.attribution import revenue as attribution_revenue
from app.domains.projects import workflow
from app.domains.reports.pdf_renderer import remove_stored_file, store_bytes
from app.domains.reports import ensure_vkpi_reports_schema

logger = get_logger(__name__)

# 单次导出硬上限。查询按 +1 取,溢出即截断并回报 truncated,避免拉爆内存/文件。
_EXPORT_ROW_LIMIT = 50000
_EXPORT_TYPES = frozenset(
    {
        "weekly",
        "projects",
        "attribution",
        "finance",
        "cost",
        "costs",
        "kpi_ledger",
        "staff_kpi",
        "kols",
        "vkpi_kol_pool",
        "favorites",
        "project_kols",
    }
)
_FINANCIAL_EXPORT_TYPES = frozenset(
    {"weekly", "attribution", "finance", "cost", "costs", "kpi_ledger", "staff_kpi"}
)
_SENSITIVE_EXPORT_FIELDS = frozenset(
    {
        "approval_note",
        "contact_links_json",
        "contact_raw_json",
        "evidence_json",
        "filters_json",
        "metadata_json",
        "note",
        "notes",
        "order_id",
        "shopify_order_snapshot_id",
        "source_ref",
        "tracking_number",
    }
)


class ExportExpired(Exception):
    """导出已过期(expires_at < now):下载端点据此返回 410 Gone,不再泄漏过期 PII/财务内容。"""


def _parse_expiry(value: Any) -> datetime | None:
    """expires_at 为 timestamptz;compat 读回可能是 datetime 或 ISO 串。解析为 UTC-aware;不可解析 → None。"""
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, str) and value.strip():
        try:
            parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError:
            return None
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    return None


def _is_expired(expires_at: Any) -> bool:
    """expires_at 缺失/不可解析 → 视为未过期(不误杀);否则 now > expires_at 即过期。"""
    if expires_at in (None, ""):
        return False
    dt = _parse_expiry(expires_at)
    if dt is None:
        return False
    return datetime.now(timezone.utc) > dt


def _cleanup_expired_file(file_path: Any) -> None:
    """惰性清理:过期导出文件 unlink(best-effort,不影响主流程);失败记日志,不静默吞。"""
    path = str(file_path or "").strip()
    if not path:
        return
    try:
        remove_stored_file(path)
    except FileNotFoundError:
        return
    except (OSError, ValueError):
        logger.warning("export expired-file cleanup failed", extra={"file_path": path}, exc_info=True)
# 电子表格公式注入前缀:以此开头的字符串单元格前置单引号,阻止 =CMD()/DDE 等被求值。
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _utcnow() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _to_usd(value: Any) -> float | None:
    """cents(整数)→ USD(两位小数)。None/非数原样为 None(导出留空)。与 PDF _money_cents 同 ÷100 口径。"""
    try:
        return round(int(value) / 100, 2)
    except (TypeError, ValueError):
        return None


def _normalize_money(row: dict[str, Any]) -> dict[str, Any]:
    """把行内所有 *_cents 列换算为 *_usd(÷100),保持列的原顺序不排序。

    kpi_ledger 的 metric_value 当 metric_key 以 _cents 结尾时同样按货币换算,
    并把 metric_key 标注为 *_usd,避免调用方误把 cents 当美元。
    """
    metric_key = row.get("metric_key")
    metric_is_money = isinstance(metric_key, str) and metric_key.endswith("_cents")
    out: dict[str, Any] = {}
    for key, value in row.items():
        if key.endswith("_cents"):
            out[key[:-6] + "_usd"] = _to_usd(value)
        elif key == "metric_value" and metric_is_money:
            out[key] = _to_usd(value)
        else:
            out[key] = value
    if metric_is_money:
        out["metric_key"] = metric_key[:-6] + "_usd"
    return out


def _guard_formula(text: str) -> str:
    return "'" + text if text[:1] in _FORMULA_PREFIXES else text


def _csv_cell(value: Any) -> Any:
    if isinstance(value, str):
        return _guard_formula(value)
    return value


def _xlsx_cell(value: Any) -> Any:
    if isinstance(value, str):
        return _guard_formula(ILLEGAL_CHARACTERS_RE.sub("", value))
    return value


def _fieldnames(rows: list[dict[str, Any]]) -> list[str]:
    """列序=首见顺序(即 SELECT 显式列序,id 等主键在前),不做字母排序。"""
    fields: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                fields.append(key)
    return fields or ["empty"]


def _uid(prefix: str) -> str:
    return f"{prefix}-{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')}-{secrets.token_hex(3)}"


def _json(value: Any) -> str:
    return json.dumps(value or {}, ensure_ascii=False, default=str)


def _staff_id(staff: dict[str, Any] | None) -> int:
    return workflow.staff_id(staff) or 0


def _rollback_export_transaction(conn: Any) -> None:
    try:
        conn.rollback()
    except Exception as exc:  # pragma: no cover - a dead connection is recovered externally
        logger.warning("vkpi export rollback failed: %s", type(exc).__name__)


def _persist_failed_export(export_id: int, exc: Exception) -> None:
    """Persist failure in a clean transaction without ever downgrading ready."""
    error_type = type(exc).__name__[:80]
    for attempt in range(2):
        recovery: Any | None = None
        try:
            recovery = get_conn()
            result = recovery.execute(
                """
                UPDATE vkpi_export_jobs
                SET status='failed', error_message=?
                WHERE id=? AND status='running'
                """,
                (error_type, int(export_id)),
            )
            recovery.commit()
            # rowcount=0 is safe: the insert may not have committed, or an
            # ambiguous commit may already have made the job ready.
            if getattr(result, "rowcount", None) == 0:
                logger.warning("vkpi export failure state not applied | export_id=%s", export_id)
            return
        except Exception as recovery_exc:
            if recovery is not None:
                _rollback_export_transaction(recovery)
            if attempt:
                logger.error(
                    "vkpi export failed-state persistence failed | export_id=%s error=%s",
                    export_id,
                    type(recovery_exc).__name__,
                )


def _cleanup_owned_export_file(stored: dict[str, Any] | None) -> None:
    if not stored or not stored.get("file_path"):
        return
    try:
        remove_stored_file(
            str(stored["file_path"]),
            expected_size=stored.get("file_size_bytes"),
            expected_sha256=str(stored.get("sha256_hex") or ""),
        )
    except FileNotFoundError:
        return
    except Exception as cleanup_exc:  # pragma: no cover - permissions/filesystem specific
        logger.error(
            "vkpi export file cleanup failed | error=%s",
            type(cleanup_exc).__name__,
        )


def _filters_with_file_integrity(
    filters: dict[str, Any],
    *,
    file_size: int,
    sha256_hex: str,
) -> dict[str, Any]:
    """Persist integrity metadata additively until dedicated columns exist."""
    return {
        **filters,
        "_file_integrity": {
            "schema_version": "export-file.v1",
            "file_size_bytes": int(file_size),
            "sha256_hex": str(sha256_hex or "").strip().lower(),
        },
    }


def _normalize_export_payload(payload: dict[str, Any] | None) -> tuple[str, dict[str, Any]]:
    raw = payload if isinstance(payload, dict) else {}
    export_type = str(
        raw.get("report_type") or raw.get("reportType") or raw.get("export_type") or "weekly"
    ).strip().lower()
    if export_type not in _EXPORT_TYPES:
        raise ValueError("unsupported export_type")
    filters: dict[str, Any] = {"report_type": export_type}
    integer_aliases = {
        "staff_id": ("staff_id", "staffId"),
        "project_id": ("project_id", "projectId"),
        "period_days": ("period_days",),
    }
    for target, aliases in integer_aliases.items():
        candidate = next((raw.get(alias) for alias in aliases if raw.get(alias) not in (None, "")), None)
        if candidate is None:
            continue
        try:
            value = int(candidate)
        except (TypeError, ValueError):
            raise ValueError(f"invalid {target}") from None
        if value > 0:
            filters[target] = min(value, 366) if target == "period_days" else value
    string_aliases = {
        "date_from": ("date_from", "startDate"),
        "date_to": ("date_to", "endDate"),
        "country": ("country",),
        "region": ("region",),
    }
    for target, aliases in string_aliases.items():
        value = next((str(raw.get(alias)).strip() for alias in aliases if raw.get(alias) not in (None, "")), "")
        if value:
            filters[target] = value[:64]
    return export_type, filters


def _assert_export_create_access(export_type: str, staff: dict[str, Any] | None) -> None:
    if not _staff_id(staff) or not check_tab_permission(staff or {}, "vkpi", "write"):
        raise scope.ScopeDenied("export requires vkpi write permission")
    if export_type == "vkpi_kol_pool" and not scope.can_view_all(staff, domain="general"):
        raise scope.ScopeDenied("full KOL pool export requires management scope")


def _is_sensitive_export_field(field: str) -> bool:
    key = str(field or "").strip().lower()
    return (
        key in _SENSITIVE_EXPORT_FIELDS
        or key == "email"
        or key == "phone"
        or key.endswith("_email")
        or key.endswith("_phone")
        or any(marker in key for marker in ("password", "secret", "token", "api_key"))
    )


def _strip_sensitive_fields(row: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in row.items() if not _is_sensitive_export_field(key)}


def _date_filter(filters: dict[str, Any], column: str) -> tuple[str, list[Any]]:
    where: list[str] = []
    params: list[Any] = []
    if filters.get("date_from"):
        where.append(f"{column}>=?")
        params.append(filters["date_from"])
    if filters.get("date_to"):
        where.append(f"{column}<=?")
        params.append(filters["date_to"])
    return (" AND " + " AND ".join(where)) if where else "", params


def _rows(export_type: str, filters: dict[str, Any], *, staff: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    conn = get_conn()
    if export_type in {"weekly", "projects"}:
        where_parts: list[str] = []
        params: list[Any] = []
        scope_clause, scope_params = scope.project_filter("p", staff, filters.get("staff_id"))
        if scope_clause:
            where_parts.append(scope_clause)
            params.extend(scope_params)
        where = f"WHERE {' AND '.join(where_parts)}" if where_parts else ""
        rows = conn.execute(
            f"""
            SELECT p.id, p.project_uid, p.project_name, p.stage, p.product_sku, p.product_name,
                   p.platform, p.marketplace, p.sample_status,
                   p.created_at, p.updated_at, k.channel_name AS kol_name,
                   u.name AS staff_name
            FROM vkpi_projects p
            LEFT JOIN kols k ON k.id = p.kol_id
            LEFT JOIN staff s ON s.id = p.assigned_staff_id
            LEFT JOIN users u ON u.id = s.user_id
            {where}
            ORDER BY p.updated_at DESC, p.id DESC
            LIMIT {_EXPORT_ROW_LIMIT + 1}
            """,
            tuple(params),
        ).fetchall()
        return [dict(row) for row in rows]
    if export_type in {"attribution", "finance"}:
        extra, params = _date_filter(filters, "sa.occurred_at")
        # 与页面 list_attributions 同源:剔除已删除项目的归因,导出行数/金额对齐页面。
        extra += f" AND {attribution_revenue._active_project_filter('sa')}"
        verified_predicate = business_truth.verified_shopify_attribution_sql("sa")
        if export_type == "finance":
            extra += f" AND {verified_predicate}"
        scope_clause, scope_params = scope.row_staff_filter("sa", staff, filters.get("staff_id"), domain="finance" if export_type == "finance" else "general")
        if scope_clause:
            extra += f" AND {scope_clause}"
            params.extend(scope_params)
        rows = conn.execute(
            f"""
            SELECT sa.id, sa.source_platform, sa.project_id, sa.link_id,
                   sa.kol_id, sa.staff_id, sa.product_sku,
                   sa.revenue_cents, sa.commission_cents, sa.currency, sa.attribution_model,
                   sa.confidence, sa.occurred_at, sa.imported_at, sa.created_at,
                   CASE WHEN {verified_predicate} THEN 1 ELSE 0 END
                       AS is_verified_business_truth,
                   p.project_name, k.channel_name AS kol_name, u.name AS staff_name
            FROM vkpi_sales_attributions sa
            LEFT JOIN vkpi_projects p ON p.id = sa.project_id
            LEFT JOIN kols k ON k.id = sa.kol_id
            LEFT JOIN staff s ON s.id = sa.staff_id
            LEFT JOIN users u ON u.id = s.user_id
            WHERE 1=1 {extra}
            ORDER BY sa.occurred_at DESC, sa.id DESC
            LIMIT {_EXPORT_ROW_LIMIT + 1}
            """,
            tuple(params),
        ).fetchall()
        output: list[dict[str, Any]] = []
        for raw_row in rows:
            row = dict(raw_row)
            verified = int(row.get("is_verified_business_truth") or 0) == 1
            row["business_truth_status"] = (
                "provider_verified" if verified else "reference_only"
            )
            if not verified:
                # Reference rows remain available to the attribution audit, but
                # cannot expose their amounts in the summable finance columns.
                row["reference_revenue_cents"] = row.get("revenue_cents")
                row["reference_commission_cents"] = row.get("commission_cents")
                row["revenue_cents"] = None
                row["commission_cents"] = None
            output.append(row)
        return output
    if export_type in {"cost", "costs"}:
        extra, params = _date_filter(filters, "cl.incurred_at")
        # Current financial exports contain only approved actual spend. Draft,
        # estimated and pending rows remain visible in their detail APIs.
        extra += f" AND {business_truth.approved_actual_cost_sql('cl')}"
        scope_clause, scope_params = scope.row_staff_filter("cl", staff, filters.get("staff_id"), domain="cost")
        if scope_clause:
            extra += f" AND {scope_clause}"
            params.extend(scope_params)
        rows = conn.execute(
            f"""
            SELECT cl.id, cl.project_id, cl.kol_id, cl.staff_id, cl.cost_type, cl.amount_cents,
                   cl.currency, cl.status, cl.incurred_at,
                   cl.created_by_staff_id, cl.created_at, cl.approved_by_staff_id, cl.approved_at,
                   cl.voided_by_staff_id, cl.voided_at, cl.updated_at,
                   'approved_actual' AS business_truth_status,
                   p.project_name, k.channel_name AS kol_name, u.name AS staff_name
            FROM vkpi_cost_ledger cl
            LEFT JOIN vkpi_projects p ON p.id = cl.project_id
            LEFT JOIN kols k ON k.id = cl.kol_id
            LEFT JOIN staff s ON s.id = cl.staff_id
            LEFT JOIN users u ON u.id = s.user_id
            WHERE 1=1 {extra}
            ORDER BY cl.incurred_at DESC, cl.id DESC
            LIMIT {_EXPORT_ROW_LIMIT + 1}
            """,
            tuple(params),
        ).fetchall()
        return [dict(row) for row in rows]
    if export_type in {"kpi_ledger", "staff_kpi"}:
        where_parts: list[str] = [business_truth.current_kpi_ledger_sql()]
        params: list[Any] = []
        scope_clause, scope_params = scope.row_staff_filter("", staff, filters.get("staff_id"))
        if scope_clause:
            where_parts.append(scope_clause)
            params.extend(scope_params)
        where = f"WHERE {' AND '.join(where_parts)}"
        rows = conn.execute(
            f"""
            SELECT id, ledger_date, staff_id, kol_id, project_id, metric_key,
                   metric_value, source_type, confidence, created_at
            FROM vkpi_kpi_ledger
            {where}
            ORDER BY ledger_date DESC, id DESC
            LIMIT {_EXPORT_ROW_LIMIT + 1}
            """,
            tuple(params),
        ).fetchall()
        return [dict(row) for row in rows]
    if export_type == "kols":
        where_parts: list[str] = []
        params: list[Any] = []
        scoped_staff_id = scope.effective_staff_id(staff, filters.get("staff_id"))
        if scoped_staff_id:
            where_parts.append("(k.assigned_staff_id=? OR k.created_by_staff_id=? OR c.staff_id=?)")
            params.extend([scoped_staff_id, scoped_staff_id, scoped_staff_id])
        where = f"WHERE {' AND '.join(where_parts)}" if where_parts else ""
        rows = conn.execute(
            f"""
            SELECT k.id, k.channel_name, k.channel_url, k.platform, k.country, k.niche,
                   k.scale_tier, k.content_type, k.primary_category, k.promoted_product,
                   k.follower_count, k.avg_views, k.contact_status,
                   k.assigned_staff_id, k.created_by_staff_id, k.created_at, k.updated_at,
                   c.status AS claim_status, c.staff_id AS claim_staff_id,
                   c.claimed_at, c.released_at
            FROM kols k
            LEFT JOIN vkpi_kol_claims c ON c.kol_id = k.id AND c.status='active'
            {where}
            ORDER BY k.id DESC
            LIMIT {_EXPORT_ROW_LIMIT + 1}
            """,
            tuple(params),
        ).fetchall()
        return [dict(row) for row in rows]
    if export_type in {"vkpi_kol_pool", "favorites", "project_kols"}:
        return _kol_pool_rows(export_type, filters, staff=staff)
    raise ValueError("unsupported export_type")


# ---- KOL 名单导出列(P0-5)。纯只读 SELECT,零触 viltrox_fit_score 写点。----
# 注:CSV/XLSX 引擎按 SELECT 显式列序出列(_fieldnames 首见顺序,不排序)。
# 联系方式不进入批量导出;明文仅允许走逐条 contact-reveal 审计路径。
# 友商关系列(competitor_brands/risk_tier)留待 P1-1 vkpi_competitor_relation 落库后再加
# (当前 0 行;且 TEXT 列 MAX 非严重度序、string_agg 跨库不可移植,P1-1 一并正确实现)。
_KOL_POOL_BASE_SELECT = """
    SELECT kp.handle AS handle,
           kp.platform AS platform,
           kp.viltrox_fit_score AS fit_score,
           kp.primary_topic AS kol_type,
           kp.country AS country,
           kp.followers AS followers,
           kp.engagement_rate AS engagement_rate
    FROM vkpi_kol_pool kp
"""


def _region_dedup_hooks(filters: dict[str, Any]) -> tuple[str, list[Any]]:
    """红线 hook:P0-6 地区过滤 +(若 P0-4 落地)duplicate_of_id IS NULL。

    本 concern 仅留参数位,实际过滤接其他 concern。默认空串=不过滤,行为与现状一致。
    其他 concern 落地时:读 filters['region']/filters['country'] 追加 'kp.country = ?',
    并视开关追加 'kp.duplicate_of_id IS NULL'(列 mig 109 已存在,实跑确认)。
    """
    extra: list[str] = []
    params: list[Any] = []
    # --- P0-6 region filter hook (no-op until wired by region concern) ---
    # --- P0-4 duplicate_of_id IS NULL hook (no-op until wired) ---
    clause = (" AND " + " AND ".join(extra)) if extra else ""
    return clause, params


def _pool_rows_gated(rows: list[Any], staff: dict[str, Any] | None) -> list[dict[str, Any]]:
    """Bulk exports never bypass the audited, per-record contact reveal flow."""
    del staff
    return [_strip_sensitive_fields(dict(row)) for row in rows]


def _kol_pool_rows(export_type: str, filters: dict[str, Any], *, staff: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    conn = get_conn()
    hook_clause, hook_params = _region_dedup_hooks(filters)
    if export_type == "vkpi_kol_pool":
        rows = conn.execute(
            f"""
            {_KOL_POOL_BASE_SELECT}
            WHERE 1=1 {hook_clause}
            ORDER BY COALESCE(kp.viltrox_fit_score, 0) DESC, kp.id DESC
            LIMIT {_EXPORT_ROW_LIMIT + 1}
            """,
            tuple(hook_params),
        ).fetchall()
        return _pool_rows_gated(rows, staff)
    if export_type == "favorites":
        # 当前员工收藏:始终强制 actor staff,view-all 也不放大到他人收藏。
        scoped_staff_id = scope.actor_staff_id(staff) or scope.effective_staff_id(staff, filters.get("staff_id"))
        if not scoped_staff_id:
            return []
        rows = conn.execute(
            f"""
            {_KOL_POOL_BASE_SELECT}
            JOIN vkpi_kol_pool_favorites f ON f.kol_pool_id = kp.id
            WHERE f.staff_id = ? {hook_clause}
            ORDER BY f.created_at DESC, f.id DESC
            LIMIT {_EXPORT_ROW_LIMIT + 1}
            """,
            (scoped_staff_id, *hook_params),
        ).fetchall()
        return _pool_rows_gated(rows, staff)
    # project_kols:某项目的 KOL 名单。
    try:
        project_id = int(filters.get("project_id") or filters.get("projectId") or 0)
    except (TypeError, ValueError):
        project_id = 0
    if not project_id:
        raise ValueError("project_kols export requires project_id")
    scope.assert_project_access(project_id, staff)
    rows = conn.execute(
        f"""
        {_KOL_POOL_BASE_SELECT}
        JOIN vkpi_project_kol_assignments a ON a.kol_pool_id = kp.id
        WHERE a.project_id = ? {hook_clause}
        ORDER BY a.id DESC
        LIMIT {_EXPORT_ROW_LIMIT + 1}
        """,
        (project_id, *hook_params),
    ).fetchall()
    return _pool_rows_gated(rows, staff)


def _csv_bytes(rows: list[dict[str, Any]]) -> bytes:
    output = io.StringIO()
    fields = _fieldnames(rows)
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    for row in rows:
        writer.writerow({key: _csv_cell(row.get(key, "")) for key in fields})
    return output.getvalue().encode("utf-8-sig")


def _xlsx_bytes(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "V-KPI Export"
    fields = _fieldnames(rows)
    ws.append([_xlsx_cell(field) for field in fields])
    for row in rows:
        ws.append([_xlsx_cell(row.get(key, "")) for key in fields])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def create_export(*, export_format: str, payload: dict[str, Any], staff: dict[str, Any] | None = None, request_meta: dict[str, str] | None = None) -> dict[str, Any]:
    ensure_vkpi_reports_schema()
    fmt = str(export_format or "").lower().strip()
    if fmt not in {"csv", "xlsx", "pdf"}:
        raise ValueError("unsupported export format")
    export_type, safe_filters = _normalize_export_payload(payload)
    _assert_export_create_access(export_type, staff)
    if fmt == "pdf" and export_type != "weekly":
        raise ValueError("pdf export supports weekly reports only")
    actor_id = _staff_id(staff)
    export_uid = _uid(f"export-{fmt}")
    conn = get_conn()
    scope.assert_legacy_default_organization(staff, conn, feature="export")
    export_id = 0
    try:
        conn.execute(
            """
            INSERT INTO vkpi_export_jobs
                (export_uid, requested_by_staff_id, export_type, file_format, filters_json, status, triggered_at, expires_at)
            VALUES (?,?,?,?,?,?,?,?)
            """,
            (export_uid, actor_id, export_type, fmt, _json(safe_filters), "running", _utcnow(), (datetime.now(timezone.utc) + timedelta(days=7)).strftime("%Y-%m-%dT%H:%M:%SZ")),
        )
        row = conn.execute("SELECT id FROM vkpi_export_jobs WHERE export_uid=?", (export_uid,)).fetchone()
        if not row:
            raise RuntimeError("export job insert was not observable")
        export_id = int(row["id"])
        conn.commit()
    except Exception:
        _rollback_export_transaction(conn)
        raise

    truncated = False
    owned_stored_file: dict[str, Any] | None = None
    try:
        if fmt == "pdf":
            report = reports.generate_weekly_report(
                period_days=int(safe_filters.get("period_days") or 7),
                staff=staff,
                filters=safe_filters,
                render_pdf=True,
            )
            file_path = str(report.get("file", {}).get("file_path") or "")
            file_size = int(report.get("file", {}).get("file_size_bytes") or 0)
            sha = str(report.get("file", {}).get("sha256_hex") or "")
            download_url = f"/api/admin/vkpi/exports/{export_id}/download"
            row_count = len(report.get("context", {}).get("projects") or [])
        else:
            rows = _rows(export_type, safe_filters, staff=staff)
            # 查询按 _EXPORT_ROW_LIMIT+1 取:超出即截断,并回报 truncated 让调用方知晓不完整。
            truncated = len(rows) > _EXPORT_ROW_LIMIT
            if truncated:
                rows = rows[:_EXPORT_ROW_LIMIT]
            rows = [_normalize_money(_strip_sensitive_fields(row)) for row in rows]
            content = _csv_bytes(rows) if fmt == "csv" else _xlsx_bytes(rows)
            stored = store_bytes(content, filename=f"{export_uid}.{fmt}")
            owned_stored_file = stored
            file_path = stored["file_path"]
            file_size = int(stored["file_size_bytes"])
            sha = stored["sha256_hex"]
            download_url = f"/api/admin/vkpi/exports/{export_id}/download"
            row_count = len(rows)
        persisted_filters = _filters_with_file_integrity(
            safe_filters,
            file_size=file_size,
            sha256_hex=sha,
        )
        ready_result = conn.execute(
            """
            UPDATE vkpi_export_jobs
            SET status='ready', file_path=?, download_url=?, row_count=?, completed_at=?,
                filters_json=?, error_message=''
            WHERE id=? AND status='running'
            """,
            (file_path, download_url, row_count, _utcnow(), _json(persisted_filters), export_id),
        )
        if getattr(ready_result, "rowcount", None) == 0:
            raise RuntimeError("export job is no longer running")
        conn.commit()
    except Exception as exc:
        # A PostgreSQL statement error poisons the transaction.  Roll back
        # before cleanup/recovery so the original exception is never masked by
        # InFailedSqlTransaction.  PDF files belong to their report run and are
        # intentionally not removed by export-job recovery.
        _rollback_export_transaction(conn)
        _cleanup_owned_export_file(owned_stored_file)
        _persist_failed_export(export_id, exc)
        raise

    audit_status = "logged"
    try:
        audit.log_export(
            staff_id=actor_id,
            export_kind="pdf_report" if fmt == "pdf" else fmt,
            export_target=export_type,
            filters=safe_filters,
            row_count=row_count,
            file_size_bytes=file_size,
            file_sha256=sha,
            download_url=download_url,
            contains_pii=export_type in {"kols", "attribution", "finance", "vkpi_kol_pool", "favorites", "project_kols"},
            contains_financial=export_type in {"weekly", "attribution", "finance", "cost", "costs"},
            ip=(request_meta or {}).get("ip", ""),
            user_agent=(request_meta or {}).get("user_agent", ""),
            metric_keys=["views", "sales", "cost", "kpi"],
        )
    except Exception as audit_exc:
        # The ready transaction is already durable.  Audit transport failure is
        # a retry/ops signal, not evidence that generation failed; never
        # downgrade or delete a valid export here.
        _rollback_export_transaction(conn)
        audit_status = "pending_retry"
        logger.error(
            "vkpi export audit pending retry | export_id=%s error=%s",
            export_id,
            type(audit_exc).__name__,
        )
    return {
        "export_id": export_id,
        "exportId": export_id,
        "export_uid": export_uid,
        "status": "ready",
        "download_url": download_url,
        "downloadUrl": download_url,
        "row_count": row_count,
        "truncated": truncated,
        "audit_status": audit_status,
    }


def _assert_export_access(item: dict[str, Any], staff: dict[str, Any] | None) -> None:
    actor_id = _staff_id(staff)
    if not actor_id:
        raise scope.ScopeDenied("export scope denied")
    requested_by = int(item.get("requested_by_staff_id") or 0)
    if requested_by == actor_id or scope.can_view_all(staff, domain="general"):
        return
    export_type = str(item.get("export_type") or "")
    if export_type in _FINANCIAL_EXPORT_TYPES and scope.can_view_all(staff, domain="finance"):
        return
    raise scope.ScopeDenied("export scope denied")


def export_file(export_id: int, *, staff: dict[str, Any] | None = None) -> dict[str, Any]:
    ensure_vkpi_reports_schema()
    conn = get_conn()
    scope.assert_legacy_default_organization(staff, conn, feature="export")
    row = conn.execute(
        """
        SELECT id, export_uid, requested_by_staff_id, export_type, file_format,
               status, file_path, download_url, download_expires_at, row_count,
               triggered_at, completed_at, expires_at, filters_json, error_message
        FROM vkpi_export_jobs
        WHERE id=?
        """,
        (int(export_id),),
    ).fetchone()
    if not row or not str(row["file_path"] or ""):
        raise LookupError("export file not found")
    item = dict(row)
    _assert_export_access(item, staff)
    if str(item.get("status") or "") != "ready":
        raise LookupError("export file is not ready")
    # 过期闸:expires_at < now → 410 Gone(路由映射),同时惰性清理落盘文件,
    # 杜绝过期导出仍可下载含 PII/财务的旧文件、以及文件无界增长。
    if _is_expired(item.get("expires_at")):
        _cleanup_expired_file(item.get("file_path"))
        raise ExportExpired("export link expired")
    try:
        persisted_filters = json.loads(str(item.get("filters_json") or "{}"))
    except (TypeError, ValueError):
        persisted_filters = {}
    integrity = persisted_filters.get("_file_integrity") if isinstance(persisted_filters, dict) else {}
    if not isinstance(integrity, dict):
        integrity = {}
    item["file_size_bytes"] = integrity.get("file_size_bytes")
    item["sha256_hex"] = str(integrity.get("sha256_hex") or "")
    item.pop("filters_json", None)
    return item


def list_exports(limit: int = 50, staff_id: int | None = None, *, staff: dict[str, Any] | None = None) -> dict[str, Any]:
    ensure_vkpi_reports_schema()
    actor_id = _staff_id(staff)
    if not actor_id:
        raise scope.ScopeDenied("export scope denied")
    conn = get_conn()
    scope.assert_legacy_default_organization(staff, conn, feature="export")
    params: list[Any] = []
    if scope.can_view_all(staff, domain="general"):
        if staff_id:
            where = "WHERE requested_by_staff_id=?"
            params.append(int(staff_id))
        else:
            where = ""
    elif scope.can_view_all(staff, domain="finance"):
        financial_types = sorted(_FINANCIAL_EXPORT_TYPES)
        placeholders = ", ".join("?" for _ in financial_types)
        if staff_id and int(staff_id) != actor_id:
            where = f"WHERE requested_by_staff_id=? AND export_type IN ({placeholders})"
            params.extend([int(staff_id), *financial_types])
        else:
            where = f"WHERE requested_by_staff_id=? OR export_type IN ({placeholders})"
            params.extend([actor_id, *financial_types])
    else:
        where = "WHERE requested_by_staff_id=?"
        params.append(actor_id)
    params.append(max(1, min(200, int(limit or 50))))
    rows = conn.execute(
        f"""
        SELECT id, export_uid, requested_by_staff_id, export_type, file_format,
               status, file_path, download_url, row_count, triggered_at, completed_at, expires_at,
               error_message
        FROM vkpi_export_jobs
        {where}
        ORDER BY triggered_at DESC, id DESC
        LIMIT ?
        """,
        tuple(params),
    ).fetchall()
    # 隐藏过期行(不给前端死链)并顺手惰性清理其落盘文件(被动 GC,控文件增长)。
    items: list[dict[str, Any]] = []
    for row in rows:
        item = dict(row)
        truth_invalidated = str(item.get("status") or "") == "invalidated"
        item["truth_invalidated"] = truth_invalidated
        # Revoked financial artifacts are immutable audit evidence.  They are
        # never exposed to clients, but also never lazily deleted by expiry GC.
        if not truth_invalidated and _is_expired(item.get("expires_at")):
            _cleanup_expired_file(item.get("file_path"))
            continue
        # Keep the retained file path and integrity record available to the
        # audit store, but never return a clickable URL for a truth-revoked
        # financial artifact.  The download endpoint independently requires
        # status=ready, so this is defence in depth for old/new clients.
        if truth_invalidated:
            item["download_url"] = ""
            item["downloadUrl"] = ""
        item.pop("file_path", None)
        items.append(item)
    return {"exports": items, "count": len(items)}
